"""Small helpers for building the guide documents with python-docx."""

from __future__ import annotations

import html as _html
import os
import re
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HEADING = RGBColor(0x00, 0x00, 0x00)
INK = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x4D, 0x4D, 0x4D)
CONTENT_WIDTH = Cm(16.6)
EDGE = Path("C:/Program Files (x86)/Microsoft/Edge/Application/msedge.exe")

_CSS = """
body { font-family: Calibri, Arial, sans-serif; font-size: 11pt; color: #1a1a1a; margin: 0 1.6cm; line-height: 1.4; }
h1 { color: #000000; font-size: 18pt; margin-top: 18pt; } h2 { color: #000000; font-size: 14pt; margin-top: 14pt; }
h3 { font-size: 12pt; margin-top: 12pt; }
.title { text-align: center; color: #000000; font-size: 24pt; font-weight: bold; margin-top: 30pt; }
.center { text-align: center; } .grey { color: #4d4d4d; }
table { border-collapse: collapse; width: 100%; margin: 6pt 0 12pt; font-size: 9.5pt; page-break-inside: avoid; }
th, td { border: 1px solid #a3a3a3; padding: 4px 6px; vertical-align: top; text-align: left; } th { background: #e5e5e5; }
.callout { background: #f2f2f2; border: 1px solid #a3a3a3; padding: 8px 10px; margin: 8pt 0 12pt; page-break-inside: avoid; }
.warn { background: #fef9c3; }
pre { background: #111111; color: #fafafa; padding: 8px 10px; font-family: Consolas, monospace; font-size: 8.8pt;
      white-space: pre-wrap; page-break-inside: avoid; }
code { font-family: Consolas, monospace; font-size: 9.5pt; } img { width: 100%; }
.cap { text-align: center; color: #4d4d4d; font-size: 9pt; margin-bottom: 12pt; }
.pb { page-break-after: always; }
"""


def _inline(text: str) -> str:
    """**bold** and `code` markup as HTML."""
    out = []
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if part.startswith("**"):
            out.append(f"<b>{_html.escape(part[2:-2])}</b>")
        elif part.startswith("`"):
            out.append(f"<code>{_html.escape(part[1:-1])}</code>")
        else:
            out.append(_html.escape(part))
    return "".join(out)


def _shade(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    props.append(shd)


class Builder:
    """A thin wrapper so guide scripts read as content, not formatting."""

    def __init__(self) -> None:
        self.doc = Document()
        for section in self.doc.sections:
            section.top_margin = section.bottom_margin = Cm(2)
            section.left_margin = section.right_margin = Cm(2.2)
        normal = self.doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(11)
        normal.paragraph_format.space_after = Pt(6)
        self.html: list[str] = []
        for name, size in (("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)):
            style = self.doc.styles[name]
            style.font.name = "Calibri"
            style.font.size = Pt(size)
            style.font.color.rgb = INK if name == "Heading 3" else HEADING

    # -- text --------------------------------------------------------------

    @staticmethod
    def runs(paragraph, text: str, size: float | None = None, colour=None):
        """Add text with **bold** and `code` markup."""
        for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
            if not part:
                continue
            if part.startswith("**"):
                run = paragraph.add_run(part[2:-2])
                run.bold = True
                if size:
                    run.font.size = Pt(size)
            elif part.startswith("`"):
                run = paragraph.add_run(part[1:-1])
                run.font.name = "Consolas"
                run.font.size = Pt((size or 11) - 1)
            else:
                run = paragraph.add_run(part)
                if size:
                    run.font.size = Pt(size)
            if colour is not None:
                run.font.color.rgb = colour
        return paragraph

    def p(self, text: str = "", size: float | None = None, colour=None, center: bool = False):
        para = self.doc.add_paragraph()
        self.runs(para, text, size, colour)
        if center:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        classes = ("center " if center else "") + ("grey" if colour is GREY else "")
        self.html.append(f'<p class="{classes.strip()}">{_inline(text)}</p>')
        return para

    def h(self, text: str, level: int = 1):
        self.html.append(f"<h{level}>{_html.escape(text)}</h{level}>")
        return self.doc.add_heading(text, level=level)

    def bullets(self, items, style: str = "List Bullet"):
        tag = "ol" if style == "List Number" else "ul"
        self.html.append(f"<{tag}>" + "".join(f"<li>{_inline(i)}</li>" for i in items) + f"</{tag}>")
        for item in items:
            self.runs(self.doc.add_paragraph(style=style), item)

    def numbered(self, items):
        self.bullets(items, style="List Number")

    def checklist(self, items):
        self.html.append("<ul style='list-style:none'>" + "".join(f"<li>&#9744; {_inline(i)}</li>" for i in items) + "</ul>")
        for item in items:
            self.runs(self.doc.add_paragraph(style="List Bullet"), "☐  " + item)

    # -- blocks ------------------------------------------------------------

    def table(self, header, rows, widths_cm, font: float = 9.5, bold_first_col: bool = False):
        self.html.append("<table><tr>" + "".join(f"<th>{_inline(h)}</th>" for h in header) + "</tr>"
                         + "".join("<tr>" + "".join(f"<td>{_inline(str(c))}</td>" for c in row) + "</tr>"
                                   for row in rows) + "</table>")
        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, text in enumerate(header):
            cell = t.rows[0].cells[i]
            cell.text = ""
            self.runs(cell.paragraphs[0], f"**{text}**", font)
            _shade(cell, "E5E5E5")
        for row in rows:
            cells = t.add_row().cells
            for i, text in enumerate(row):
                cells[i].text = ""
                self.runs(cells[i].paragraphs[0], f"**{text}**" if bold_first_col and i == 0 else str(text), font)
        for row in t.rows:
            for i, width in enumerate(widths_cm):
                row.cells[i].width = Cm(width)
        self.doc.add_paragraph()
        return t

    def callout(self, text: str, fill: str = "F2F2F2"):
        css = "callout warn" if fill.upper() == "FEF9C3" else "callout"
        self.html.append(f'<div class="{css}">' + "<br>".join(_inline(line) for line in text.split("\n")) + "</div>")
        t = self.doc.add_table(rows=1, cols=1)
        t.style = "Table Grid"
        cell = t.rows[0].cells[0]
        cell.text = ""
        for i, line in enumerate(text.split("\n")):
            para = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            self.runs(para, line, 10.5)
        _shade(cell, fill)
        cell.width = CONTENT_WIDTH
        self.doc.add_paragraph()

    def code(self, text: str):
        """A dark block for commands or code; one line of text per line of input."""
        self.html.append(f"<pre>{_html.escape(text)}</pre>")
        t = self.doc.add_table(rows=1, cols=1)
        cell = t.rows[0].cells[0]
        cell.text = ""
        for i, line in enumerate(text.split("\n")):
            para = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            para.paragraph_format.space_after = Pt(0)
            run = para.add_run(line if line else " ")
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0xF8, 0xFA, 0xFC)
        _shade(cell, "111111")
        cell.width = CONTENT_WIDTH
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def image(self, path: Path, caption: str, width_cm: float = 16.6):
        self.html.append(f'<img src="{Path(path).resolve().as_uri()}" style="width: min({width_cm}cm, 100%); '
                         f'display: block; margin: 0 auto"><div class="cap">{_inline(caption)}</div>')
        self.doc.add_picture(str(path), width=Cm(width_cm))
        self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = self.doc.add_paragraph()
        self.runs(cap, caption, 9, GREY)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(12)

    def page_break(self):
        self.html.append('<div class="pb"></div>')
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def title(self, text: str, subtitle: str, lines: list[str]):
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(text)
        run.bold = True
        run.font.size = Pt(24)
        run.font.color.rgb = HEADING
        self.html.append(f'<div class="title">{_html.escape(text)}</div>')
        self.p(subtitle, size=13, colour=INK, center=True)
        for line in lines:
            self.p(line, size=10.5, colour=GREY, center=True)

    def save(self, path: Path):
        self.doc.save(path)

    def save_pdf(self, path: Path):
        """Write the same content as HTML and print it to PDF with Microsoft Edge (no extra software)."""
        html_path = Path(path).with_suffix(".html")
        page = ("<!doctype html><html><head><meta charset='utf-8'><style>" + _CSS + "</style></head><body>"
                + "\n".join(self.html) + "</body></html>")
        html_path.write_text(page, encoding="utf-8")
        subprocess.run([str(EDGE), "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                        f"--print-to-pdf={Path(path).resolve()}", html_path.resolve().as_uri()],
                       check=True, capture_output=True, timeout=300)
        if not os.environ.get("KEEP_HTML"):
            html_path.unlink()
